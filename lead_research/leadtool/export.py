"""leads.xlsx (ready rows first, wide wrapped email_body) and SUMMARY.md."""
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from . import icypeas, store

ORDER = {"ready": 0, "needs_check": 1, "skip": 2}
WIDTHS = {"email_body": 90, "why_fit": 50, "research_note": 50, "subject": 40,
          "guessed_unverified": 40, "contact_source_url": 40}


def _rank(r):
    return ORDER.get(r["status"].split(":")[0], 3)


def export():
    rows = sorted(store.load_leads(), key=_rank)
    wb = Workbook()
    ws = wb.active
    ws.title = "leads"
    ws.append(store.COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([r.get(c, "") for c in store.COLUMNS])
    for i, col in enumerate(store.COLUMNS, 1):
        letter = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[letter].width = WIDTHS.get(col, 20)
        if col == "email_body":
            for cell in ws[letter][1:]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    wb.save(os.path.join(store.ROOT, "leads.xlsx"))

    q = [r for r in rows if store.is_qualified(r)]
    ready = [r for r in q if r["status"] == "ready"]
    generic = [r for r in ready if r["email_source"].startswith("generic")]
    personal = [r for r in ready if not r["email_source"].startswith("generic")]
    needs = [r for r in q if r["status"].startswith("needs_check")]
    skipped = [r for r in rows if not store.is_qualified(r)]
    by_country = {}
    for r in q:
        by_country[r["country"]] = by_country.get(r["country"], 0) + 1
    lines = [
        "# Lead research summary",
        "",
        f"- Qualified agencies found: {len(q)}",
        f"- Ready to send: {len(ready)}",
        f"  - to a personal address of the decision-maker: {len(personal)}",
        f"  - generic address only (owner named in the greeting): {len(generic)}",
        f"- Needs check: {len(needs)}",
        f"- Researched and skipped (did not qualify): {len(skipped)}",
        f"- Icypeas credits used: {icypeas.credits_used()}",
        "",
        "By country: " + ", ".join(f"{k or 'unknown'} {v}" for k, v in sorted(by_country.items())),
        "",
        "See DECISIONS.md for how the run was done and its limits.",
    ]
    with open(os.path.join(store.ROOT, "SUMMARY.md"), "w") as f:
        f.write("\n".join(lines) + "\n")
    print("\n".join(lines))
